"""Daane 2021 Data S1 (mmc2.xlsx): the supplement's own lat4a / slc43a2a naming, an Ensembl ortholog link to the zebrafish gene, corrected sheet descriptions,
and append-only records. Earlier files (daane2021_pins.json, lat4a_resolution_v2.json, the 16 + 24 records) are NOT edited: they describe what was known before this was read.

  PYTHONPATH=. python3 -m connectors.genomics.daane2021_supplement

The workbook is fetched and hash-checked, not stored. The four Ensembl REST responses in data/genomics/artifacts were fetched with
  https://rest.ensembl.org/homology/id/oryzias_latipes/<ID>?content-type=application/json;target_species=danio_rerio;type=orthologues;format=condensed
  https://rest.ensembl.org/lookup/id/<ID>?content-type=application/json
and are read back here by hash.
"""

from __future__ import annotations

import hashlib
import io
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "genomics"
ART = OUT / "artifacts"
XLSX_URL = "https://ars.els-cdn.com/content/image/1-s2.0-S0960982221011908-mmc2.xlsx"
XLSX_SHA = "1ac98650f493036218837d235ed4451ff096d926fa50b797dcea4fc7fca52279"
LAT4A, LAT4B = "ENSORLG00000009715", "ENSORLG00000005019"          # medaka (Oryzias latipes) Ensembl ids as printed in the workbook
NS, ISSUER, CREATED_AT = "node-a", "genomics-worker", "2026-10-31T23:30:00+00:00"
# What each sheet is, from its own header row and sheet name (read 2026-10-31), replacing the earlier guess in daane2021_pins.json.
SHEETS = {"A - fin measurement": "per-specimen fin and standard-length measurements", "B - sample info": "species, specimen and locality for sequenced samples",
          "C - cov. GO enrich.": "GO enrichment by capture coverage", "D - heterozygosity": "heterozygous and fixed species-specific SNPs per species",
          "E - CNV": "average exon copy number and predicted duplicate genes per species", "F - gliders RER enrich.": "enrichment for relative evolutionary rate, gliding vs other beloniforms",
          "G - conv. aa subst.": "convergent amino-acid substitutions (gene, SNP, species columns, fin-screen and fin/limb GO flags)", "H - fin and limb genes": "fin and limb candidate genes with average coverage",
          "I - Glider RER genes": "per-gene relative-evolutionary-rate rho, p and q for gliders", "J - Exocoeidae phylop": "per-gene phyloP acceleration or conservation (cons/acc) with p for the Exocoetidae",
          "K - accel CNEs Exocoet.": "GO terms for accelerated conserved noncoding elements, Exocoetidae", "L - con CNEs Exocoet.": "GO terms for constrained CNEs, Exocoetidae",
          "M - accel CNEs halfbeak": "GO terms for accelerated CNEs, halfbeaks", "N - con CNEs halfbeak": "GO terms for constrained CNEs, halfbeaks"}


def read_rows():
    import openpyxl
    sys.path.insert(0, str(ROOT))
    from connectors.genomics import sources as S
    S.ALLOWED_HOSTS = S.ALLOWED_HOSTS | {"ars.els-cdn.com"}
    raw, rc = S.fetch(XLSX_URL, 5_000_000)
    if rc.sha256 != XLSX_SHA:
        raise SystemExit("Data S1 workbook bytes changed: " + rc.sha256)
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    assert list(wb.sheetnames) == list(SHEETS), wb.sheetnames
    want = {}
    for sh in ("G - conv. aa subst.", "H - fin and limb genes", "I - Glider RER genes", "J - Exocoeidae phylop"):
        rows = list(wb[sh].iter_rows(values_only=True))
        want[sh] = {"header": [c for c in rows[0] if c is not None], "rows": [{"sheet_row": i + 1, "cells": [c for c in r if c is not None]} for i, r in enumerate(rows)
                                                                              if any(c in (LAT4A, LAT4B) for c in r)]}
    return want, rc


def main() -> int:
    sys.path.insert(0, str(ROOT))
    from frontier import models, store
    rows, rc = read_rows()
    sha = lambda p: hashlib.sha256((ART / p).read_bytes()).hexdigest()
    ens = {k: json.loads((ART / f).read_text()) for k, f in (("hom_a", f"ensembl_homology_{LAT4A}.json"), ("hom_b", f"ensembl_homology_{LAT4B}.json"), ("lk_a", f"ensembl_lookup_{LAT4A}.json"), ("lk_b", f"ensembl_lookup_{LAT4B}.json"))}
    zf = {r["symbol"]: r for r in __import__("csv").DictReader(open(OUT / "genes.csv"))}
    hom = lambda h: [x["id"] for x in h["data"][0]["homologies"] if x["type"] == "ortholog_one2one"]
    assert hom(ens["hom_a"]) == [zf["slc43a2a"]["ensembl_gene"]] and hom(ens["hom_b"]) == [zf["slc43a2b"]["ensembl_gene"]] and ens["hom_a"]["data"][0]["id"] == LAT4A
    G, H, I, J = (rows[k]["rows"] for k in rows)
    cell = lambda rs, rid: next(r["cells"] for r in rs if rid in r["cells"])
    assert cell(G, LAT4A)[1:3] == ["lat4a", "F326L"] and cell(H, LAT4A)[1] == "lat4a" and cell(H, LAT4B)[1] == "lat4b" and cell(I, LAT4B)[1] == "slc43a2b"
    assert cell(J, LAT4A)[1:] == ["slc43a2a", "acc", 0.00837] and cell(J, LAT4B)[1:] == ["slc43a2b", "acc", 1e-05]
    prev = OUT / "daane2021_supplement_sheets.json"      # keep the first witnessed retrieval time so a rebuild does not change record versions
    first = (json.loads(prev.read_text())["workbook"].get("retrieved_at") if prev.exists() else None) or rc.retrieved_at
    wbmeta = {"url": XLSX_URL, "sha256": XLSX_SHA, "bytes": rc.bytes, "stored_in_repo": False, "retrieved_at": first}
    (OUT / "daane2021_supplement_sheets.json").write_text(json.dumps({
        "schema": "daane2021-supplement-sheets-0.1 (connector-local)", "workbook": wbmeta,
        "supersedes_for_descriptions": "daane2021_pins.json supplement.items['Data S1C-S1N'], whose per-letter descriptions were guessed from how the main text cites them and are partly wrong (E is copy number, not contigs; J is per-gene phyloP, K-N are the CNE tables). daane2021_pins.json is left unchanged as the record of that earlier state.",
        "basis": "each sheet's own name and header row, read from the workbook", "sheets": SHEETS, "lat4a_rows": rows}, indent=1, sort_keys=True))
    chain = [
        {"step": 1, "what": "the paper's convergent F326L gene is labelled lat4a in Data S1G and S1H", "where": "Data S1G row %d, S1H row %d" % (G[0]["sheet_row"], next(r["sheet_row"] for r in H if LAT4A in r["cells"])), "id": LAT4A, "id_space": "medaka Ensembl (Oryzias latipes)"},
        {"step": 2, "what": "the same medaka id is labelled slc43a2a in Data S1J (phyloP acceleration p = 0.00837; the main text rounds it to 0.008)", "where": "Data S1J row %d" % next(r["sheet_row"] for r in J if LAT4A in r["cells"]), "id": LAT4A},
        {"step": 3, "what": "Ensembl lists one zebrafish ortholog (ortholog_one2one, Clupeocephala) for that medaka gene", "id": zf["slc43a2a"]["ensembl_gene"], "artifact": f"ensembl_homology_{LAT4A}.json", "sha256": sha(f"ensembl_homology_{LAT4A}.json")},
        {"step": 4, "what": "that zebrafish Ensembl id is ZFIN slc43a2a (id_crosscheck.json, genes.csv)", "id": zf["slc43a2a"]["zfin_id"], "zfin_chr": zf["slc43a2a"]["zfin_gff_chr"]},
        {"step": 5, "what": "consistency: the paper maps nr21 to chromosome 15, where ZFIN puts slc43a2a (slc43a2b is on chromosome 10)", "where": "daane2021_pins.json quote lat4a_map"}]
    control = {"what": "lat4b is a different medaka gene (%s; Data S1H lat4b, S1I and S1J slc43a2b) and Ensembl maps it to a different zebrafish gene" % LAT4B, "zebrafish_ensembl": zf["slc43a2b"]["ensembl_gene"], "zfin_id": zf["slc43a2b"]["zfin_id"],
               "distinct_from_lat4a": zf["slc43a2b"]["ensembl_gene"] != zf["slc43a2a"]["ensembl_gene"], "artifact": f"ensembl_homology_{LAT4B}.json", "sha256": sha(f"ensembl_homology_{LAT4B}.json")}
    res = {"schema": "lat4a-resolution-0.3 (connector-local)", "supersedes": "lat4a_resolution_v2.json, kept unchanged as the record of the state before the supplement was read",
           "question": "Which ZFIN gene is Daane et al. 2021's lat4a?",
           "status": "SUPPORTED: slc43a2a (ZDB-GENE-040426-964, ENSDARG00000036848) by the supplement's own gene naming plus an Ensembl one-to-one ortholog; the paper's text and the workbook state no ZFIN accession",
           "chain": chain, "negative_control": control, "workbook": wbmeta,
           "still_not_established": ["the authors do not print a zebrafish accession for lat4a; the link from their medaka-annotation name to the zebrafish gene is Ensembl's orthology call, not theirs",
                                     "the residue check (T200K, F326L) is still inconclusive for slc43a2a: its RefSeq protein NP_956545.1 is 153 residues (see lat4a_resolution_v2.json), so the paper's numbering was not re-derived here",
                                     "the mapping interval in Figures S2D-S2F was not read",
                                     "UniProt still names the slc43a2b product with the LAT4 name; that is a protein-name annotation and is kept as a disagreement, not resolved"],
           "reading": "Use slc43a2a for 'lat4a', labelled 'supported by the supplement and Ensembl orthology; not an author-stated zebrafish accession'. The earlier 'probable, not confirmed' wording is replaced by this, not deleted."}
    (OUT / "lat4a_resolution_v3.json").write_text(json.dumps(res, indent=1, sort_keys=True))
    # ---- records (append-only; predecessor None; the superseded records keep their versions) ----
    old = json.loads((OUT / "daane2021_records.json").read_text())["records"]
    oldv = {r["id"]: r for r in old}
    def sup(wb_loc):
        return {"uri": XLSX_URL, "locator": wb_loc, "source_version": "Current Biology 2021, Data S1 (mmc2.xlsx), sha256 pinned", "content_sha256": XLSX_SHA, "content_available": False,
                "access": "publisher CDN download, no login; read level: full-text for sheets G, H, I, J rows naming lat4a and slc43a2b only; workbook not stored; licence not read", "license": None}
    def ens_src(loc, art):
        return {"uri": "https://rest.ensembl.org/", "locator": loc, "source_version": "Ensembl REST (release not stated in the response)", "content_sha256": sha(art), "content_available": False,
                "access": "public REST GET, no key; response stored by hash; Ensembl terms not read", "license": None}
    tag = lambda cls, sp: f"[species: {sp}; evidence_class: {cls}; source: Daane 2021 Data S1 + Ensembl REST] "
    def claim(id_, title, st, ctx, lim, source, origin, assumptions=None):
        return {"namespace": NS, "kind": "Claim", "id": id_, "issuer": ISSUER, "source_uri": "urn:frontier:connector:genomics", "visibility": "public", "predecessor": None,
                "payload": {"title": title, "statement": st, "context": ctx, "assumptions": assumptions or [], "limitations": lim, "source": source, "origin": origin}}
    def aref(path, title, role):
        b = (ROOT / path).read_bytes(); h = hashlib.sha256(b).hexdigest()
        return {"title": f"{role}: {path} - {title}"[:3000], "uri": "urn:sha256:" + h, "sha256": h, "byte_size": len(b), "media_type": "application/json", "role": role, "availability": "referenced", "license": None}
    def evid(id_, target, rel, source, method, context, arts, lim, indep, when):
        return {"namespace": NS, "kind": "Evidence", "id": id_, "issuer": ISSUER, "source_uri": "urn:frontier:connector:genomics", "visibility": "public", "predecessor": None,
                "payload": {"target": target, "evidence_kind": "source_assessment", "relation": rel, "source": source, "created_time": {"value": when, "basis": "witnessed"}, "retrieved_time": {"value": when, "basis": "witnessed"},
                            "method": method, "context": context, "assumptions": [], "provenance": {"workbook_sha256": XLSX_SHA, "ensembl_artifacts": {k: sha(f) for k, f in (("homology_lat4a", f"ensembl_homology_{LAT4A}.json"), ("homology_lat4b", f"ensembl_homology_{LAT4B}.json"))}},
                            "source_dependence": {"declared": True, "independent_of_author": indep, "note": "one operator read the files and wrote the records; no independent review"},
                            "uncertainty": "Not quantified.", "limitations": lim, "artifacts": [aref(p, t, r) for p, t, r in arts]}}
    sheets_art = ("data/genomics/daane2021_supplement_sheets.json", "the exact lat4a / lat4b rows with sheet and row numbers, and corrected sheet descriptions", "output")
    v3_art = ("data/genomics/lat4a_resolution_v3.json", "the five-step chain, its negative control and what is still not established", "output")
    D = [claim("gx-daane-datas1-names-lat4a-as-slc43a2a", "Source statement (Data S1G, S1H, S1J): the supplement labels medaka ENSORLG00000009715 lat4a and slc43a2a",
               "Data S1G lists ENSORLG00000009715 as lat4a with the substitution F326L; S1H lists the same id as lat4a among fin and limb genes (average coverage 0.92); S1J lists the same id as slc43a2a with phyloP 'acc' p = 0.00837. "
               "ENSORLG00000005019 is lat4b in S1H and slc43a2b in S1I and S1J. The ids are medaka Ensembl ids, as printed in the workbook.",
               tag("source_assessment", "Oryzias latipes (ids), Beloniformes") + "Read in the workbook's own cells; exact rows are in daane2021_supplement_sheets.json.",
               ["These are the authors' gene labels in medaka coordinates, not zebrafish identifiers.", "Only the rows naming lat4a, lat4b and slc43a2b were read; the other sheets were not examined beyond their headers.", "The workbook is not stored; its licence was not read."],
               sup("Data S1 workbook (mmc2.xlsx), sheets G, H, I, J"), "source_statement"),
         claim("gx-lat4a-is-slc43a2a-supported", "Inference (supported, not author-stated): the paper's lat4a is zebrafish slc43a2a (ZDB-GENE-040426-964, ENSDARG00000036848)",
               "The supplement names the paper's lat4a gene slc43a2a (Data S1J); Ensembl gives that medaka gene one zebrafish ortholog, ENSDARG00000036848, which is ZFIN slc43a2a on chromosome 15, where the paper maps nr21. The control holds: lat4b (slc43a2b, "
               "ENSORLG00000005019) maps to the different zebrafish gene ENSDARG00000061120 (ZFIN slc43a2b, chromosome 10). This replaces the status of gx-lat4a-is-probably-slc43a2a, which is kept unchanged.",
               tag("source_assessment", "Danio rerio") + "Chain and control in lat4a_resolution_v3.json.",
               ["The authors print no zebrafish accession; the medaka-to-zebrafish link is Ensembl's one-to-one ortholog call.", "The residue check (T200K, F326L) is still inconclusive for slc43a2a (RefSeq protein of 153 residues).", "The mapping interval in Figures S2D-S2F was not read.", "UniProt's LAT4 protein name on the slc43a2b product disagrees and is kept as a disagreement."],
               ens_src("homology/id/oryzias_latipes/%s (target danio_rerio)" % LAT4A, f"ensembl_homology_{LAT4A}.json"), "derived", ["Ensembl's ortholog_one2one call for the medaka gene is correct."])]
    P = [{"namespace": NS, "kind": "Problem", "id": "gx-problem-lat4a-residues", "issuer": ISSUER, "source_uri": "urn:frontier:connector:genomics", "visibility": "public", "predecessor": None,
          "payload": {"title": "Re-derive the paper's T200K and F326L positions on a full-length zebrafish slc43a2a protein", "known_state": "The identity of lat4a as slc43a2a is now supported by the supplement's naming and an Ensembl ortholog (gx-lat4a-is-slc43a2a-supported). What is not done: slc43a2a's RefSeq protein is 153 residues, too short to contain residues 200 or 326, so the paper's numbering was never reproduced.",
                      "question": "Does a full-length slc43a2a (for example the Ensembl ENSDARP00000120945 translation) carry T at the paper's position 200 and the paper's F at 326 in its own numbering?", "problem_kind": "project_gap", "novelty_status": "Not searched beyond the paper, the supplement rows and the stored sources.",
                      "data_access": "Data S1 rows G, H, I, J read; the 19-page PDF read in full; Figures S1-S4 and Tables S1-S3 not retrieved.", "completeness": "recorded", "required_capabilities": ["the Ensembl protein sequence for ENSDARP00000120945", "the paper's Figure 3C alignment"],
                      "downstream_uses": ["Would close the last 'inferred' step on the residue level."], "dependencies": [], "source": sup("Data S1 workbook (mmc2.xlsx), sheets G and J"),
                      "original": {"supersedes_for_status": "gx-problem-lat4a-accession (kept unchanged; records are append-only)"}, "approach": "unknown", "scales": ["molecular"], "context_tags": ["species:Danio rerio", "evidence_class:source_assessment"]}}]
    import datetime
    ens_when = datetime.datetime.fromtimestamp(max((ART / f"ensembl_homology_{g}.json").stat().st_mtime for g in (LAT4A, LAT4B)), datetime.timezone.utc).isoformat()   # when the stored responses were written
    made = {d["id"]: models.make_record(models.Draft.model_validate(d), created_at=CREATED_AT) for d in D + P}
    E = [evid("gx-ev-datas1-lat4a-rows", "gx-daane-datas1-names-lat4a-as-slc43a2a", "supports", sup("Data S1 workbook sheets G, H, I, J"), "Fetched the workbook, checked its sha256 against the pinned value, and read the rows that contain the two medaka ids.",
              "The authors' own file; exact rows and sheet numbers are in the artifact.", [sheets_art], ["Only the named rows were read.", "The labels are the authors', in medaka coordinates."], False, first),
         evid("gx-ev-ensembl-ortholog", "gx-lat4a-is-slc43a2a-supported", "supports", ens_src("homology/id/oryzias_latipes/%s" % LAT4A, f"ensembl_homology_{LAT4A}.json"),
              "Queried Ensembl for the zebrafish orthologs of the two medaka genes, stored the four responses by hash, and compared the zebrafish ids with ZFIN's own table (genes.csv, id_crosscheck.json); lat4b is the negative control.",
              "An independent database, but it is an orthology call and not the authors' statement.", [v3_art, sheets_art, ("data/genomics/id_crosscheck.json", "ZFIN and Ensembl ids for the two zebrafish genes", "input"),
              (f"data/genomics/artifacts/ensembl_homology_{LAT4A}.json", "Ensembl homology response for the lat4a medaka gene", "input"), (f"data/genomics/artifacts/ensembl_homology_{LAT4B}.json", "Ensembl homology response for the lat4b control", "input")],
              ["Ensembl release not stated in the response.", "The residue check remains inconclusive.", "Orthology is not the authors' accession."], True, ens_when)]
    for e in E:
        e["payload"]["target"] = models.ref(made[e["payload"]["target"]]); made[e["id"]] = models.make_record(models.Draft.model_validate(e), created_at=CREATED_AT)
    recs = list(made.values())
    bad = store._unlicensed_available_content(recs)
    if bad:
        raise SystemExit(f"licence gate: {bad}")
    (OUT / "daane2021_supplement_records.json").write_text(json.dumps({"records_schema": "genomics-records-0.1 (connector-local; frontier.models records)", "status": "proposed drafts; nothing submitted, reviewed or accepted", "namespace": NS, "issuer": ISSUER,
        "append_only_note": "adds 5 records. gx-lat4a-is-probably-slc43a2a and gx-problem-lat4a-accession in daane2021_records.json keep their versions and are superseded in status by gx-lat4a-is-slc43a2a-supported and gx-problem-lat4a-residues, not edited. "
                            "Their texts say 'not confirmed' and 'the supplement was not retrieved', which was true when written.",
        "supersedes_in_status": {"gx-lat4a-is-probably-slc43a2a": oldv["gx-lat4a-is-probably-slc43a2a"]["version"] if "version" in oldv["gx-lat4a-is-probably-slc43a2a"] else None}, "records": recs}, indent=1, sort_keys=True))
    print({k: sum(1 for r in recs if r["kind"] == k) for k in ("Claim", "Evidence", "Problem")}, "licence offenders:", len(bad))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
